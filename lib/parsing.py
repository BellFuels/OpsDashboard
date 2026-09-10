"""
parsing.py — reads the Bell_Unified_<date>.xlsx workbook entirely from memory.

Helpers are copied from scripts/build_unified.py (the file's producer) so the
consumer and producer agree on every format quirk. Nothing here ever touches
disk: load_unified() takes raw bytes and returns DataFrames.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, time as dt_time, timedelta

import openpyxl
import pandas as pd

SCHEMA_VERSION = "2"
EXCEL_EPOCH_OFFSET = 25569  # days between 1899-12-30 and 1970-01-01

DRIVER_SENIORITY = ["Jeff", "Augustine", "Christopher", "Raul", "Pascual",
                    "Eric", "Bino", "Vicente", "Dan", "Brett", "Mataeo"]
ALL_SERVICE_TYPES = ["FLEET", "GEN", "TANK", "TANK/SHOW", "GRVTY", "GRVTY/PUMP"]

DELIVERY_COLUMNS = ["Date", "Driver", "Stop", "SO", "Product", "Gallons", "StopMins",
                    "Units", "Address", "FleetType", "CustType", "GPM",
                    "Arrival", "Departure", "IsFleet", "IsTerminal"]
PAYROLL_COLUMNS = ["Date", "Driver", "Hours", "ClockIn", "ClockOut",
                   "BackToYard", "DowntimeStart", "DowntimeEnd", "DowntimeNote"]
PUNCH_COLUMNS = ["Date", "Driver", "Seq", "In", "Out"]
CUSTOMER_COLUMNS = ["Name", "Account", "CustType", "SvcType", "Street", "City", "County", "FullAddress"]


class UnifiedFileError(Exception):
    """Raised when the uploaded file isn't a valid unified workbook."""


# ─── Cell utilities (ports of build_unified.py / v1 JS) ──────────────────────

def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_float(v):
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def parse_int(v):
    f = parse_float(v)
    return int(f) if f is not None else 0


def serial_to_datetime(n):
    return datetime(1970, 1, 1) + timedelta(days=n - EXCEL_EPOCH_OFFSET)


def extract_date_iso(v):
    """Date cell -> ISO YYYY-MM-DD. Accepts datetimes, Excel serials, ISO text,
    M/D/YYYY text — defensive against Excel re-saves converting text dates."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)):
        try:
            return serial_to_datetime(float(v)).strftime("%Y-%m-%d")
        except (OverflowError, ValueError):
            return ""
    s = str(v).strip()
    if not s:
        return ""
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    part = s.split(" ")[0]
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})$", part)
    if m:
        mo, da, yr = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if yr < 100:
            yr += 2000
        try:
            return datetime(yr, mo, da).strftime("%Y-%m-%d")
        except ValueError:
            return ""
    try:
        return serial_to_datetime(float(s)).strftime("%Y-%m-%d")
    except ValueError:
        return ""


def arrival_to_datetime(v):
    """Arrival/Departure cell -> pd.Timestamp (NaT if unparseable).
    Handles 'YYYY-MM-DD HH:MM:SS' text, datetime objects, and Excel serials."""
    if v is None or v == "":
        return pd.NaT
    if isinstance(v, datetime):
        return pd.Timestamp(v)
    if isinstance(v, (int, float)):
        if float(v) < 40000:
            return pd.NaT
        return pd.Timestamp(serial_to_datetime(float(v))).round("s")
    s = str(v).strip()
    if not s:
        return pd.NaT
    if re.match(r"^\d{4,6}(\.\d+)?$", s):
        return arrival_to_datetime(float(s))
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%m/%d/%Y %I:%M:%S %p",
                "%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%Y-%m-%d"):
        try:
            return pd.Timestamp(datetime.strptime(s, fmt))
        except ValueError:
            continue
    return pd.NaT


def parse_hhmmss(v):
    """'HH:MM:SS' / 'MM:SS' / decimal minutes -> decimal minutes (None if invalid)."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    parts = s.split(":")
    try:
        if len(parts) == 3:
            return int(parts[0]) * 60 + int(parts[1]) + int(parts[2]) / 60
        if len(parts) == 2:
            return int(parts[0]) + int(parts[1]) / 60
        return float(s)
    except ValueError:
        return None


def fmt_hmm(mins):
    if mins is None or pd.isna(mins):
        return "—"
    mins = float(mins)
    return f"{int(mins // 60)}:{int(round(mins % 60)):02d}"


def fmt_hhmmss(mins):
    if mins is None or pd.isna(mins):
        return "—"
    total = int(round(abs(float(mins)) * 60))
    return f"{total // 3600:02d}:{(total % 3600) // 60:02d}:{total % 60:02d}"


def normalize_service_type(v):
    if not v:
        return ""
    s = re.sub(r"\s+", "", str(v).strip().upper())
    if s in ("TANK",):
        return "TANK"
    if s in ("TANK/SHOW", "TANKSHOW"):
        return "TANK/SHOW"
    if s == "FLEET":
        return "FLEET"
    if s == "GEN":
        return "GEN"
    if s in ("GRVTY/PUMP", "GRVTYPUMP"):
        return "GRVTY/PUMP"
    if s == "GRVTY":
        return "GRVTY"
    return str(v).strip().upper()


# ─── Workbook loading ────────────────────────────────────────────────────────

@dataclass
class UnifiedData:
    deliveries: pd.DataFrame
    payroll: pd.DataFrame
    punches: pd.DataFrame
    customers: pd.DataFrame
    meta: dict
    product_map: dict
    benchmarks: dict          # svcType -> minutes (float)
    shift_split_time: str     # "HH:MM"
    threshold: int
    driver_order: list
    deliveries_no_fleet: pd.DataFrame = field(default=None)
    rolled_history: pd.DataFrame = field(default=None)
    averages: pd.DataFrame = field(default=None)

    @property
    def dates(self):
        return sorted(self.deliveries["date"].unique())


def _sheet_lists(wb, name):
    if name not in wb.sheetnames:
        return []
    rows = []
    it = wb[name].iter_rows(values_only=True)
    next(it, None)  # header
    for raw in it:
        if raw is None or not any(c is not None and cell_str(c) for c in raw):
            continue
        rows.append(list(raw))
    return rows


def clock_str(v):
    """Manual time cell -> 'H:MM AM/PM' text. Excel stores typed times as real
    time values (the file may be uploaded before a build normalizes them), so
    accept datetimes, times, day-fractions, and 12/24-hour text."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        v = v.time()
    if isinstance(v, dt_time):
        return f"{v.hour % 12 or 12}:{v.minute:02d} {'AM' if v.hour < 12 else 'PM'}"
    if isinstance(v, (int, float)) and 0 <= v < 1:
        mins = int(round(v * 1440)) % 1440
        h, mm = divmod(mins, 60)
        return f"{h % 12 or 12}:{mm:02d} {'AM' if h < 12 else 'PM'}"
    s = str(v).strip()
    m = re.match(r"^(\d{1,2}):(\d{2})(?::\d{2})?\s*(AM|PM)?$", s, re.I)
    if not m:
        return s
    h, mm, ap = int(m.group(1)), int(m.group(2)), m.group(3)
    if ap:
        h = h % 12 + (12 if ap.upper() == "PM" else 0)
    return f"{h % 12 or 12}:{mm:02d} {'AM' if h % 24 < 12 else 'PM'}"


COUNTY_SUFFIX_RE = re.compile(r",\s*[A-Za-z][A-Za-z ]*$")


def normalize_stop(v):
    """Canonical stop name — collapse whitespace and case so the same customer at
    the same address keys to one row ('TRUCK KING' vs 'Truck King')."""
    return re.sub(r"\s+", " ", str(v or "").strip()).upper()


def normalize_address(a):
    """Canonical form of a delivery address, used for both display and grouping.

    The builder appends ', <COUNTY>' to any address it can match to a customer
    record, but only began doing so partway through the history — so one site
    shows up as both '5645 W 31ST STREET' and '5645 W 31ST STREET, COOK'. Case
    drifts too ('513 Express Center Dr'). Since averages key on (address, stop),
    that split a site's history into separate rows: 29 stops covering 19% of all
    deliveries, which skewed the Daily tab's vs-Avg columns and made established
    sites look like they had no baseline on the timeline.

    A trailing segment containing digits (', STE 1950', ', P-107') is a real part
    of the address and is kept.
    """
    s = re.sub(r"\s+", " ", str(a or "").strip())
    return COUNTY_SUFFIX_RE.sub("", s).upper()


_STATE_ZIP_RE = re.compile(r"\s+[A-Z]{2}\s+(\d{5})(?:-\d{4})?\s*$")
# words that can never be part of a town name (street types, unit markers, directionals)
_STREET_WORDS = frozenset({
    "STREET", "ST", "AVENUE", "AVE", "ROAD", "RD", "DRIVE", "DR", "BOULEVARD", "BLVD",
    "LANE", "LN", "COURT", "CT", "WAY", "PLACE", "PL", "PARKWAY", "PKWY", "PLAZA",
    "HIGHWAY", "HWY", "TRAIL", "SUITE", "STE", "UNIT", "APT", "BLDG", "FLOOR", "FL",
    "DOCK", "ROOM", "RM", "ATTN", "N", "S", "E", "W", "NORTH", "SOUTH", "EAST", "WEST",
})


def _street_key(s):
    """Loose comparison key for a street: drop street-type words and punctuation."""
    s = str(s or "").lower().split(",")[0]
    s = re.sub(r"\b(street|st|avenue|ave|road|rd|drive|dr|boulevard|blvd|lane|ln|"
               r"court|ct|way|place|pl|north|south|east|west)\b", " ", s)
    return re.sub(r"[^a-z0-9]", "", s)


def _streets_roughly_match(addr_street, cust_street):
    """Copied from build_unified.streets_roughly_match so producer and consumer
    agree on what counts as the same street."""
    a, b = _street_key(addr_street), _street_key(cust_street)
    if not a or not b:
        return False
    if a == b:
        return True
    if len(a) >= 6 and a in b:
        return True
    if len(b) >= 6 and b in a:
        return True
    m_a = re.match(r"^\s*(\d+)", str(addr_street))
    m_b = re.match(r"^\s*(\d+)", str(cust_street))
    if m_a and m_b and m_a.group(1) == m_b.group(1):
        rest_a, rest_b = a[len(m_a.group(1)):], b[len(m_b.group(1)):]
        if len(rest_a) >= 4 and len(rest_b) >= 4 and (rest_b[:6] in rest_a or rest_a[:6] in rest_b):
            return True
    return False


def _clean_town(tokens, known):
    toks = [t for t in tokens if t not in _STREET_WORDS and not re.fullmatch(r"[\d#.\-()/&']+", t)]
    if not toks or len(toks) > 3:
        return None
    # a leading street fragment sometimes survives ("147TH OAK FOREST"); if the
    # tail is a town we already recognise, keep only that
    for n in (2, 1):
        if len(toks) > n and " ".join(toks[-n:]) in known:
            return " ".join(toks[-n:])
    return " ".join(toks)


def build_town_lookup(customer_streets, addresses):
    """{delivery address -> town}, derived from the customer list.

    The Customers sheet's City column is empty in practice, so the town is only
    available buried inside the Street string ("5645 W 31ST STREET CICERO IL
    60804"), with no delimiter between street and town. Two passes:

      1. Where a ZIP has several customer addresses, the town is the trailing run
         of words all of them share — street parts differ, the town does not.
      2. A ZIP seen only once (158 of 423 here, e.g. Cicero) has nothing to agree
         with, so the matched customer street's prefix is stripped using the
         delivery address and what remains is the town.

    Addresses that cannot be resolved confidently are simply left out; the caller
    shows no town rather than a guessed one.
    """
    streets = [s for s in (normalize_address(x) for x in customer_streets)
               if s and _STATE_ZIP_RE.search(s)]
    if not streets:
        return {}

    by_zip = {}
    for s in streets:
        m = _STATE_ZIP_RE.search(s)
        by_zip.setdefault(m.group(1), []).append(s[:m.start()].strip().split())
    zip_town = {}
    for z, bodies in by_zip.items():
        if len(bodies) < 2:
            continue
        common = ()
        for n in (1, 2, 3):
            tails = {tuple(b[-n:]) for b in bodies if len(b) > n}
            if len(tails) == 1:
                common = next(iter(tails))
            else:
                break
        town = list(common)
        while town and town[0] in _STREET_WORDS:
            town.pop(0)
        if town:
            zip_town[z] = " ".join(town)
    known = set(zip_town.values())

    # bucket customer streets by house number so matching stays a short scan
    buckets = {}
    for s in streets:
        m = _STATE_ZIP_RE.search(s)
        body = s[:m.start()].strip()
        num = re.match(r"\s*(\d+)", body)
        buckets.setdefault(num.group(1) if num else "", []).append((body, zip_town.get(m.group(1))))

    out = {}
    for addr in addresses:
        if not addr:
            continue
        num = re.match(r"\s*(\d+)", addr)
        for body, town in buckets.get(num.group(1) if num else "", []):
            if not _streets_roughly_match(addr, body):
                continue
            if town:
                out[addr] = town
            else:
                a_toks, b_toks = addr.split(), body.split()
                i = 0
                while i < len(a_toks) and i < len(b_toks) and a_toks[i] == b_toks[i]:
                    i += 1
                guess = _clean_town(b_toks[i:], known)
                if guess:
                    out[addr] = guess
            break
    return out


def load_unified(file_bytes: bytes) -> UnifiedData:
    """Parse the unified workbook from raw bytes. Never touches disk."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        raise UnifiedFileError("Couldn't read that file as an Excel workbook — "
                               "upload the Bell_Unified_….xlsx from the daily email.")
    try:
        if "Deliveries" not in wb.sheetnames or "Meta" not in wb.sheetnames:
            raise UnifiedFileError("This doesn't look like a unified file — "
                                   "upload the Bell_Unified_….xlsx from the daily email.")
        meta = {}
        for raw in _sheet_lists(wb, "Meta"):
            if cell_str(raw[0]):
                meta[cell_str(raw[0])] = cell_str(raw[1] if len(raw) > 1 else "")
        if meta.get("schema_version") != SCHEMA_VERSION:
            raise UnifiedFileError(f'Unsupported unified file version '
                                   f'"{meta.get("schema_version", "?")}" — expected {SCHEMA_VERSION}.')

        # Deliveries
        drows = []
        for raw in _sheet_lists(wb, "Deliveries"):
            def col(i):
                return raw[i] if i < len(raw) else None
            gallons = parse_float(col(5))
            date = extract_date_iso(col(0))
            if gallons is None or not date:
                continue
            stop_mins = parse_float(col(6))
            gpm = parse_float(col(11))
            drows.append({
                "date": date, "driver": cell_str(col(1)), "stop": normalize_stop(col(2)),
                "so": cell_str(col(3)), "product": cell_str(col(4)), "gallons": gallons,
                "stop_mins": round(stop_mins) if stop_mins is not None else None,
                "units": parse_int(col(7)), "address": normalize_address(col(8)),
                "fleet_type": cell_str(col(9)), "cust_type": cell_str(col(10)),
                "gpm": gpm, "arrival": arrival_to_datetime(col(12)),
                "departure": arrival_to_datetime(col(13)),
                "is_fleet": cell_str(col(14)) == "Y", "is_terminal": cell_str(col(15)) == "Y",
            })
        if not drows:
            raise UnifiedFileError("The unified file has no delivery rows.")
        deliveries = pd.DataFrame(drows)

        # Payroll
        prows = []
        for raw in _sheet_lists(wb, "Payroll"):
            def col(i):
                return raw[i] if i < len(raw) else None
            date = extract_date_iso(col(0))
            hours = parse_float(col(2))
            if not date or not cell_str(col(1)) or hours is None:
                continue
            prows.append({"date": date, "driver": cell_str(col(1)), "hours": hours,
                          "clock_in": cell_str(col(3)), "clock_out": cell_str(col(4)),
                          "back_to_yard": clock_str(col(5)),
                          "downtime_start": clock_str(col(6)),
                          "downtime_end": clock_str(col(7)),
                          "downtime_note": cell_str(col(8))})
        payroll = pd.DataFrame(prows, columns=["date", "driver", "hours", "clock_in",
                                               "clock_out", "back_to_yard",
                                               "downtime_start", "downtime_end",
                                               "downtime_note"])

        # Punches
        purows = []
        for raw in _sheet_lists(wb, "Punches"):
            def col(i):
                return raw[i] if i < len(raw) else None
            date = extract_date_iso(col(0))
            if not date or not cell_str(col(1)):
                continue
            purows.append({"date": date, "driver": cell_str(col(1)), "seq": parse_int(col(2)),
                           "in": cell_str(col(3)), "out": cell_str(col(4))})
        punches = pd.DataFrame(purows, columns=["date", "driver", "seq", "in", "out"])
        if len(punches):
            punches = punches.sort_values(["date", "driver", "seq"]).reset_index(drop=True)

        # Customers (display info only — enrichment is baked into Deliveries)
        crows = []
        cust_streets = []
        for raw in _sheet_lists(wb, "Customers"):
            def col(i):
                return raw[i] if i < len(raw) else None
            if not cell_str(col(0)):
                continue
            crows.append({"name": cell_str(col(0)), "cust_type": cell_str(col(2)),
                          "svc_type": cell_str(col(3))})
            # col 4 is Street; it is the only place the town appears
            cust_streets.append(cell_str(col(4)))
        customers = pd.DataFrame(crows, columns=["name", "cust_type", "svc_type"])
        town_by_addr = build_town_lookup(cust_streets, deliveries["address"].unique())
        deliveries["town"] = deliveries["address"].map(town_by_addr).fillna("")
    finally:
        wb.close()

    product_map = {k[len("product."):]: v for k, v in meta.items() if k.startswith("product.")}
    benchmarks = {}
    for k, v in meta.items():
        if k.startswith("benchmark."):
            mins = parse_hhmmss(v)
            if mins is not None and mins > 0:
                benchmarks[k[len("benchmark."):].upper()] = mins

    sst = re.match(r"^(\d{1,2}):(\d{2})", meta.get("shift_split_time", "") or "")
    shift_split = f"{int(sst.group(1)):02d}:{sst.group(2)}" if sst else "12:00"

    try:
        threshold = int(meta.get("threshold", "20"))
    except ValueError:
        threshold = 20

    order = [d.strip() for d in (meta.get("driver_order") or "").split(",") if d.strip()]
    driver_order = order or list(DRIVER_SENIORITY)

    data = UnifiedData(deliveries=deliveries, payroll=payroll, punches=punches,
                       customers=customers, meta=meta, product_map=product_map,
                       benchmarks=benchmarks, shift_split_time=shift_split,
                       threshold=threshold, driver_order=driver_order)
    data.deliveries_no_fleet = deliveries[~deliveries["is_fleet"] & ~deliveries["is_terminal"]].reset_index(drop=True)

    # rolled history + averages are computed once here (calc imports parsing, not vice versa)
    from lib.calc import roll_up_stops, build_averages
    data.rolled_history = roll_up_stops(data.deliveries_no_fleet)
    data.averages = build_averages(data.rolled_history)
    return data
